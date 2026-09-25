def test_algorithm_update_deck_renders_four_nonblank_slides(tmp_path):
    pdfium = __import__("pypdfium2")
    from pptx import Presentation
    from scripts.figures.make_algorithm_update_deck import build

    build(tmp_path)
    pdf = tmp_path / "algorithm_update_2026-09-25.pdf"
    pptx = tmp_path / "algorithm_update_2026-09-25.pptx"
    with pdfium.PdfDocument(pdf) as document:
        assert len(document) == 4
        for page in document:
            text = page.get_textpage().get_text_range()
            assert len(text) > 150
    assert len(Presentation(pptx).slides) == 4
