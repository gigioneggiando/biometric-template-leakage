"""Build a concise PDF/PPTX briefing for the source-analysis contribution."""
from __future__ import annotations

from io import BytesIO
import json
from pathlib import Path

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
from matplotlib.backends.backend_pdf import PdfPages
from matplotlib.patches import FancyArrowPatch, Rectangle
from pptx import Presentation
from pptx.util import Inches


ROOT = Path(__file__).resolve().parents[2]
WIDTH, HEIGHT = 13.333, 7.5
INK = "#202725"
MUTED = "#61706a"
GREEN = "#177245"
BLUE = "#2f6690"
RED = "#b33a3a"
AMBER = "#b7791f"
PALE = "#edf2ef"


def canvas(title: str, number: int):
    figure = plt.figure(figsize=(WIDTH, HEIGHT), facecolor="white")
    figure.text(.055, .925, title, fontsize=24, weight="bold", color=INK, va="top")
    figure.text(.055, .865, "BIOMETRIC KEY-PROVENANCE ANALYSIS  |  25 SEPTEMBER 2026",
                fontsize=8, color=MUTED, va="top")
    figure.text(.055, .045, "Internal research update | Not a security certificate",
                fontsize=8, color=MUTED)
    figure.text(.945, .045, f"{number}/4", fontsize=8, color=MUTED, ha="right")
    return figure


def box(figure, x, y, w, h, title, body, color=BLUE):
    figure.patches.append(Rectangle((x, y), w, h, transform=figure.transFigure,
                                    facecolor="white", edgecolor=color, linewidth=1.5))
    figure.text(x + .018, y + h - .035, title, fontsize=13, weight="bold", color=color, va="top")
    figure.text(x + .018, y + h - .09, body, fontsize=10, color=INK, va="top", linespacing=1.45)


def arrow(figure, start, end):
    figure.patches.append(FancyArrowPatch(start, end, transform=figure.transFigure,
                                          arrowstyle="-|>", mutation_scale=15,
                                          linewidth=1.4, color=MUTED))


def slide_architecture():
    figure = canvas("What the proposed algorithm does", 1)
    figure.text(.055, .79,
                "Question: can source analysis find recurring biometric transforms, explain the path, and trigger a testable remediation?",
                fontsize=13, color=INK, va="top")
    positions = [.055, .295, .535, .775]
    content = [
        ("Python integration", "BioHash\nIoM-GRP\nPolyProtect", BLUE),
        ("V3 interpreter", "Key provenance\nControl-flow joins\nFail-closed unknowns", GREEN),
        ("Evidence", "Sink + line\nCall trace\nReuse capacity", AMBER),
        ("Separate experiment", "Fresh-key candidate\nLeakage test\nTAR / FMR gate", RED),
    ]
    for x, (title, body, color) in zip(positions, content):
        box(figure, x, .39, .185, .25, title, body, color)
    for left, right in zip(positions, positions[1:]):
        arrow(figure, (left + .188, .515), (right - .005, .515))
    figure.text(.055, .28, "Claim boundary", fontsize=13, weight="bold", color=INK)
    figure.text(.055, .235,
                "The analyzer classifies key scope under explicit contracts. It does not certify secrecy, unlinkability, deployment security, or utility.",
                fontsize=11, color=INK)
    return figure


def slide_analyzer_results():
    figure = canvas("Analyzer evidence improved from v1 to v3", 2)
    labels = ["V1 audited", "V2 development", "V2 confirmation", "V3 + real recipes", "Local baseline"]
    coverage = [69.57, 95.65, 95.0, 96.15, 23.08]
    colors = [AMBER, BLUE, GREEN, GREEN, MUTED]
    axis = figure.add_axes([.08, .22, .56, .57])
    positions = list(range(len(labels)))
    bars = axis.barh(positions, coverage, color=colors, height=.56)
    axis.set_yticks(positions, labels)
    axis.invert_yaxis()
    axis.set_xlim(0, 100)
    axis.set_xlabel("Coverage (%)")
    axis.axvline(75, color=RED, linestyle="--", linewidth=1.2, label="Frozen 75% gate")
    axis.grid(axis="x", color="#d8dedb", linewidth=.6)
    axis.set_axisbelow(True)
    axis.spines[["top", "right", "left"]].set_visible(False)
    for bar, value in zip(bars, coverage):
        axis.text(value + 1, bar.get_y() + bar.get_height() / 2, f"{value:.2f}", va="center", fontsize=9)
    axis.legend(loc="lower right", frameon=False)
    box(figure, .70, .53, .24, .22, "Frozen confirmation", "19 / 19 decisions correct\n1 conservative abstention\n0 false-fresh", GREEN)
    box(figure, .70, .25, .24, .22, "Combined v3 study", "25 / 25 decisions correct\n96.15% coverage\nAll fixed gates pass", BLUE)
    return figure


def slide_integration():
    figure = canvas("Executable scheme integration and baseline", 3)
    axis = figure.add_axes([.055, .35, .56, .43])
    axis.axis("off")
    table = axis.table(
        cellText=[
            ["BioHash", "Reuse", "Fresh"],
            ["IoM-GRP", "Reuse", "Fresh"],
            ["PolyProtect", "Reuse", "Fresh"],
        ],
        colLabels=["Executable recipe", "Recurring pool", "Record-specific"],
        cellLoc="left", colLoc="left", bbox=[0, 0, 1, 1], colWidths=[.40, .30, .30],
    )
    table.auto_set_font_size(False)
    table.set_fontsize(11)
    for (row, _), cell in table.get_celld().items():
        cell.set_edgecolor("#cdd6d1")
        cell.set_linewidth(.7)
        if row == 0:
            cell.set_facecolor(PALE)
            cell.set_text_props(weight="bold")
    box(figure, .67, .56, .27, .20, "V3", "25 decisions\n25 correct\n1 abstention", GREEN)
    box(figure, .67, .31, .27, .20, "Intraprocedural baseline", "6 decisions\n5 correct\n20 abstentions", AMBER)
    figure.text(.055, .245, "Why the gap?", fontsize=13, weight="bold", color=INK)
    figure.text(.055, .195,
                "V3 follows helpers and scheme contracts. The baseline sees one function only and misclassifies record_id // 1 as reuse.",
                fontsize=11, color=INK)
    figure.text(.055, .14,
                "This is a transparent engineering baseline, not a state-of-the-art static-analysis comparison.",
                fontsize=10, color=MUTED)
    return figure


def slide_utility():
    figure = canvas("The remaining blocker is authentication utility", 4)
    effects = json.loads((ROOT / "experiments/scface_utility_confirmation_2026-09-25/execution_manifest.json").read_text())
    if effects["status"] != "completed" or effects["completed_evaluations"] != 96:
        raise ValueError("Completed utility confirmation required")
    import csv
    with (ROOT / "experiments/scface_utility_confirmation_2026-09-25/effects.csv").open(newline="") as handle:
        rows = list(csv.DictReader(handle))
    labels = [f"{row['dataset']}\n{row['split_seed']}" for row in rows]
    means = [100 * float(row["tar_difference"]) for row in rows]
    lower = [100 * float(row["tar_lower"]) for row in rows]
    upper = [100 * float(row["tar_upper"]) for row in rows]
    axis = figure.add_axes([.08, .25, .56, .56])
    x = list(range(4))
    axis.errorbar(x, means, yerr=[[mean - lo for mean, lo in zip(means, lower)],
                                  [hi - mean for hi, mean in zip(upper, means)]],
                  fmt="o", color=BLUE, ecolor=BLUE, capsize=5, markersize=7)
    axis.axhline(-3, color=RED, linestyle="--", linewidth=1.4, label="-3 point margin")
    axis.axhline(0, color=MUTED, linewidth=.8)
    axis.set_xticks(x, labels)
    axis.set_ylabel("Fresh minus recurring TAR (points)")
    axis.grid(axis="y", color="#d8dedb", linewidth=.6)
    axis.set_axisbelow(True)
    axis.spines[["top", "right"]].set_visible(False)
    axis.legend(frameon=False, loc="upper left")
    box(figure, .70, .55, .24, .21, "Frozen result", "0 / 4 utility cells pass\nAll FMR bounds pass\nTAR bounds fail", RED)
    box(figure, .70, .29, .24, .20, "Decision for Sani", "Review novelty\nApprove utility margin\nChoose next cohort", AMBER)
    figure.text(.70, .20, "Do not claim lower risk without utility cost.", fontsize=11, weight="bold", color=RED)
    return figure


def build(destination: Path) -> None:
    destination.mkdir(parents=True, exist_ok=True)
    slides = [slide_architecture(), slide_analyzer_results(), slide_integration(), slide_utility()]
    pdf_path = destination / "algorithm_update_2026-09-25.pdf"
    pptx_path = destination / "algorithm_update_2026-09-25.pptx"
    presentation = Presentation()
    presentation.slide_width = Inches(WIDTH)
    presentation.slide_height = Inches(HEIGHT)
    presentation.core_properties.title = "Biometric key-provenance analysis update"
    with PdfPages(pdf_path) as pdf:
        for figure in slides:
            pdf.savefig(figure, bbox_inches=None)
            buffer = BytesIO()
            figure.savefig(buffer, format="png", dpi=160, facecolor="white")
            buffer.seek(0)
            slide = presentation.slides.add_slide(presentation.slide_layouts[6])
            slide.shapes.add_picture(buffer, 0, 0, width=presentation.slide_width, height=presentation.slide_height)
            plt.close(figure)
    presentation.save(pptx_path)
    print(f"Wrote {pdf_path}")
    print(f"Wrote {pptx_path}")


if __name__ == "__main__":
    build(ROOT / "reports/slides")
