"""Build an editable review deck and matching PDF without biometric images."""
from __future__ import annotations

from pathlib import Path
import argparse
from io import BytesIO

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
from matplotlib.backends.backend_pdf import PdfPages
from PIL import Image
import pandas as pd
import pypdfium2 as pdfium
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.util import Inches, Pt

ROOT = Path(__file__).resolve().parents[2]
FIGURES = ROOT / "reports/figures"
WIDTH, HEIGHT = 13.333333, 7.5
FONT = "Times New Roman"


def add_text(slide, fig, text, x, y, width, height, size=18, bold=False, color="#222222"):
    if not text.isascii():
        raise ValueError(f"Non-ASCII slide label: {text}")
    shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(width), Inches(height))
    frame = shape.text_frame
    frame.margin_left = frame.margin_right = frame.margin_top = frame.margin_bottom = 0
    frame.word_wrap = False
    for index, line in enumerate(text.split("\n")):
        paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
        paragraph.text = line
        paragraph.font.name = FONT
        paragraph.font.size = Pt(size)
        paragraph.font.bold = bold
        paragraph.font.color.rgb = RGBColor.from_string(color.lstrip("#"))
        paragraph.space_after = Pt(0)
        paragraph.line_spacing = 1.4
    artist = fig.text(x / WIDTH, 1 - y / HEIGHT, text, fontsize=size, family=FONT,
                      weight="bold" if bold else "normal", color=color, va="top", linespacing=1.4)
    fig.canvas.draw()
    bound = artist.get_window_extent(fig.canvas.get_renderer())
    region = matplotlib.transforms.Bbox.from_bounds(x * fig.dpi, (HEIGHT - y - height) * fig.dpi,
                                                    width * fig.dpi, height * fig.dpi).padded(0.5)
    if not region.contains(bound.x0, bound.y0) or not region.contains(bound.x1, bound.y1):
        raise ValueError(f"Slide text exceeds its region: {text}")


def add_figure(slide, fig, name, x=0.6, y=1.05, width=12.1, height=5.35):
    path = FIGURES / f"{name}.png"
    with Image.open(path) as image:
        ratio = image.width / image.height
    image_width = min(width, height * ratio)
    image_height = image_width / ratio
    left = x + (width - image_width) / 2
    top = y + (height - image_height) / 2
    slide.shapes.add_picture(str(path), Inches(left), Inches(top), width=Inches(image_width), height=Inches(image_height))
    ax = fig.add_axes([left / WIDTH, 1 - (top + image_height) / HEIGHT, image_width / WIDTH, image_height / HEIGHT])
    ax.imshow(plt.imread(path))
    ax.axis("off")


def build_presentation(out: Path) -> None:
    table = pd.read_csv(ROOT / "experiments/cross_dataset_key_pool_summary.csv")
    out.mkdir(parents=True, exist_ok=True)
    presentation = Presentation()
    presentation.slide_width, presentation.slide_height = Inches(WIDTH), Inches(HEIGHT)
    presentation.core_properties.title = "Hidden keys and repeated biometric records"
    presentation.core_properties.subject = "Research review draft, 2026-09-19"
    with PdfPages(out / "research_review.pdf") as pdf:
        for number, title in enumerate([
            "Hidden keys and repeated biometric records",
            "Experimental architecture",
            "Dataset and protection coverage",
            "Set reconstruction and gallery linkage",
            "Earlier three-seed key-pool studies",
            "Mechanism and correlation controls",
            "Replicated multi-record amplification",
            "Native linkage and radial sensitivity",
            "Genuine raw-input controls",
            "Failure analysis: aggregation is not always better",
            "Four datasets, distinct evidence layers",
            "Independent pools and a matched baseline",
            "Extended exposures: MOBIO and SCface",
            "Learned attacks on raw embeddings",
            "Local stricter PolyProtect selection",
        ], start=1):
            slide = presentation.slides.add_slide(presentation.slide_layouts[6])
            fig = plt.figure(figsize=(WIDTH, HEIGHT), dpi=120, facecolor="white")
            add_text(slide, fig, title, 0.6, 0.3, 12.1, 0.6, size=25, bold=True)
            add_text(slide, fig, f"Research review draft | 19 September 2026 | {number}/15", 0.6, 7.07, 12.1, 0.3, size=11, color="#666666")
            if number == 1:
                add_text(slide, fig, "Can multiple protected records reveal identity when keys remain hidden?",
                         0.6, 1.45, 12.1, 0.7, size=20)
                add_text(slide, fig, "Fresh-key learned attacks: uncertainty remains; no equivalence claim.\n\n"
                         "Recurring transforms: large gains from multiple same-identity records.\n\n"
                         "Pool draws matter; a simple prediction-mean baseline is competitive.",
                         0.6, 2.55, 12.1, 3.1, size=19)
                add_text(slide, fig, "New: 672 MOBIO/SCface endpoints, raw-input retraining and a local stricter-selection audit.",
                         0.6, 6.3, 12.1, 0.45, size=13)
            elif number in (2, 4, 5, 6):
                name, caption = {
                    2: ("fig_architecture", "Train on separate identities. Reconstruct an embedding, then link to a held-out gallery."),
                    4: ("fig_attack_detail", "Two aggregation paths; shared training objective; held-out identity linkage. Key values remain hidden."),
                    5: ("fig_results_overview", f"{len(table)} conditions from {table['source_file'].nunique()} earlier studies. New one-seed pilots are reported separately."),
                    6: ("fig_controls", "Slot identifiers are not key values. Coarse and fine correlation sweeps use separate partitions."),
                }[number]
                add_figure(slide, fig, name)
                add_text(slide, fig, caption, 0.6, 6.58, 12.1, 0.4, size=12)
            elif number == 3:
                rows = [
                    ["Dataset", "Identities", "Train/val/test", "Embeddings", "Variation"],
                    ["MOBIO", "150", "90/30/30", "1,799 / 1,800", "Sessions"],
                    ["LFW", "125", "75/25/25", "1,500 / 1,500", "In-the-wild"],
                    ["FEI", "200", "120/40/40", "2,378 / 2,400", "Pose/expression"],
                    ["SCface", "130", "78/26/26", "2,851 / 2,860", "Camera/distance"],
                ]
                for row_index, row in enumerate(rows):
                    for column, value in enumerate(row):
                        add_text(slide, fig, value, [0.6, 2.6, 4.6, 7.1, 10.0][column], 1.5 + row_index * 0.5,
                                 [1.8, 1.8, 2.3, 2.7, 2.7][column], 0.5, size=15, bold=row_index == 0)
                add_text(slide, fig, "BioHash: 128 bits; four datasets; includes a Haar-sign control on MOBIO.\n"
                         "MLP-Hash: 512 bits; MOBIO; paper-specified, not source-exact.\n"
                         "IoM-GRP: 300 codes (q=16); PolyProtect: 170 reals (m=5, overlap=2).\n"
                         "Both: MOBIO/FEI original follow-up; MOBIO/SCface extended follow-up.",
                         0.6, 4.15, 12.1, 1.85, size=16)
                add_text(slide, fig, "SCface: mugshot gallery and visible surveillance probes; 9 detection failures, all identities eligible.",
                         0.6, 6.3, 12.1, 0.45, size=13)
            elif number == 7:
                add_figure(slide, fig, "fig_followup_amplification")
                add_text(slide, fig, "MOBIO / FEI; IoM-GRP / PolyProtect; 3 model seeds x 2 identity partitions; matched 120-epoch caps.",
                         0.6, 6.58, 12.1, 0.4, size=12)
            elif number == 8:
                add_figure(slide, fig, "fig_followup_native_controls")
                add_text(slide, fig, "Protected-gallery matching is distinct from learned linkage. Radial sensitivity does not establish natural norm leakage.",
                         0.6, 6.58, 12.1, 0.4, size=12)
            elif number == 9:
                add_figure(slide, fig, "fig_native_norm_audit")
                add_text(slide, fig, "Earlier native audit, not learned retraining. New raw learned and local stricter-selection results appear on slides 14-15.",
                         0.6, 6.58, 12.1, 0.4, size=12)
            elif number == 10:
                add_figure(slide, fig, "fig_followup_failures")
                add_text(slide, fig, "Original primary family unchanged. All 48 contrasts, including significant losses, are reported separately from historical exploration.",
                         0.6, 6.58, 12.1, 0.4, size=12)
            elif number == 11:
                add_figure(slide, fig, "fig_dataset_coverage")
                add_text(slide, fig, "SCface now has a three-seed/two-partition exposure study. LFW remains historical; missing cells are not zero accuracy.",
                         0.6, 6.58, 12.1, 0.4, size=12)
            elif number == 12:
                add_figure(slide, fig, "fig_pool_replication")
                add_text(slide, fig, "Three new pool draws: IoM gains persist; PolyProtect is variable. Input-pooling superiority is not established.",
                         0.6, 6.58, 12.1, 0.4, size=12)
            else:
                name, caption = {
                    13: ("fig_extended_exposures", "672 endpoints; eight primary gains pass Holm correction. Larger pools do not guarantee lower observed linkage."),
                    14: ("fig_raw_learned", "Separate descriptive study: bars are seed SD, not confidence intervals. No matched raw-versus-unit causal test."),
                    15: ("fig_stricter_selection", "No corrected reduction from this local selection rule; not a source-exact refutation of the original policy."),
                }[number]
                add_figure(slide, fig, name)
                add_text(slide, fig, caption, 0.6, 6.58, 12.1, 0.4, size=12)
            slide.notes_slide.notes_text_frame.text = (
                "Evidence baseline: commit 655ae1ecc2c9fc0aa80c828fe0303753296f66df. "
                "FEI run recorded base commit d1e4ceb3f975fb47d9a8321c47484bb1411f5239 with FEI additions uncommitted.\n"
                "Sources: experiments/cross_dataset_key_pool_summary.csv; experiments/fei_multiexposure/README.md; "
                "experiments/mobio_mechanism_controls/results_summary.csv; experiments/mobio_correlation_controls/results_summary.csv.\n"
                "Configuration: configs/attacks/fei_key_pool_boundary.yaml. See reports/figures/README.md for figure captions "
                "and docs/ROADMAP.md for pending gates. Diagram symbols are schematic, not biometric examples.\n"
                "New pilot code freeze: d5f4e89. Configuration: configs/attacks/scheme_extension_pilot.yaml. "
                "Sources: experiments/scheme_extension_pilot/{results_summary,paired_uncertainty,equivalence_sensitivity,native_utility}.csv. "
                "SCface protocol and pilot freeze: 69a93e4; sources under experiments/scface_multiexposure and "
                "experiments/scface_scheme_extension_pilot. "
                "One-seed CPU pilots, 120-epoch cap; not a controlled ranking against earlier 400-epoch, three-seed studies. "
                "See figure_appendix.pdf for all figures, including native matching and uncertainty."
                " Follow-up: experiments/scheme_followup_2026-09-18; 216 endpoints, two identity partitions, three model seeds. "
                "Pre-execution source/config hashes: execution_manifest.json; base commit 4352eeb, dirty working tree recorded. "
                "Primary sign-flip tests: Holm family 8; native gallery-label permutations: Holm family 12."
                " Independent-pool follow-up: experiments/pool_replication_2026-09-19; 24 cells, 144 fits, "
                "72 prediction-mean evaluations; source snapshots and input hashes recorded. Three-pool intervals "
                "are pointwise, not corrected significance tests. Full-text comparison uses the author's thesis section 5.4."
            )
            pdf.savefig(fig)
            plt.close(fig)
    presentation.save(out / "research_review.pptx")
    with pdfium.PdfDocument.new() as appendix:
        for source in sorted(FIGURES.glob("fig_*.pdf")):
            with pdfium.PdfDocument(source) as document:
                appendix.import_pages(document)
        appendix.save(out / "figure_appendix.pdf")
    print(f"15-slide review deck, PDF, and complete figure appendix -> {out}")


def build_dataset_update(destination: Path) -> None:
    from pypdf import PdfReader, PdfWriter, Transformation

    table = pd.read_csv(ROOT / "experiments/cross_dataset_key_pool_summary.csv")
    scface = table[table["dataset"] == "SCface"].set_index("pool_size")
    pilot_dirs = [ROOT / "experiments" / name for name in
                  ("scheme_extension_pilot", "scface_scheme_extension_pilot")]
    pilots = pd.concat([pd.read_csv(folder / "results_summary.csv") for folder in pilot_dirs])
    paired = pd.concat([pd.read_csv(folder / "paired_uncertainty.csv") for folder in pilot_dirs])
    gains = paired[(paired["condition"] == "random_key_pool_4") &
                   (paired["model"] == "mean_mlp") & (paired["contrast"] == "ten_minus_one")]
    native = pd.concat([pd.read_csv(folder / "native_utility.csv") for folder in pilot_dirs])
    fresh_native = native[(native["scheme"] == "PolyProtect") &
                          (native["condition"] == "independent_unseen_keys")].set_index("dataset")
    followup = ROOT / "experiments/scheme_followup_2026-09-18"
    contrasts = pd.read_csv(followup / "seed_identity_contrasts.csv")
    primary = contrasts[contrasts["primary"]]
    native_controls = pd.read_csv(followup / "native_null_controls.csv")
    endpoints = pd.read_csv(followup / "seed_identity_endpoints.csv")
    model_runs = pd.read_csv(followup / "results_summary.csv")
    titles = ["September dataset and experiment update", "Experimental architecture",
              "Detailed attacker: aggregation and reconstruction", "Cross-dataset key-reuse boundary",
              "Mechanism controls and interpretation", "Multi-seed, multi-partition validation",
              "Native linkage and radial sensitivity", "Earlier cross-dataset scheme pilots",
              "Closest research and our precise contribution", "Attacker access: assumptions and limits",
              "Separate implementation and matching checks", "Raw embeddings: scale is not identity leakage",
              "Failure analysis: all 48 contrasts",
              "Findings and study scope", "Four datasets: coverage and evidence strength",
              "Independent pool draws and matched baseline", "Extended exposures: MOBIO and SCface",
              "Learned raw-input attacks: a separate study", "Local stricter PolyProtect selection"]
    figure_names = {2: "fig_architecture", 3: "fig_attack_detail", 4: "fig_results_overview", 5: "fig_controls",
                    6: "fig_followup_amplification", 7: "fig_followup_native_controls", 8: "fig_scheme_pilots",
                    12: "fig_native_norm_audit", 13: "fig_followup_failures",
                    15: "fig_dataset_coverage", 16: "fig_pool_replication", 17: "fig_extended_exposures",
                    18: "fig_raw_learned", 19: "fig_stricter_selection"}
    presentation = Presentation()
    writer = PdfWriter()
    for number, title in enumerate(titles, start=1):
        slide = presentation.slides.add_slide(presentation.slide_layouts[6])
        fig = plt.figure(figsize=(WIDTH, HEIGHT), dpi=120, facecolor="white")
        add_text(slide, fig, title, 0.6, 0.3, 12.1, 0.6, size=25, bold=True)
        add_text(slide, fig, f"19 September 2026 | Earlier: 4352eeb / 15e4384; latest results through 4831d99; manifests retained | {number}/{len(titles)}",
                 0.6, 7.07, 12.1, 0.3, size=11, color="#666666")
        if number == 1:
            add_text(slide, fig, "Identity-disjoint evaluation across session, pose and camera variation.",
                     0.6, 1.15, 12.1, 0.5, size=17)
            rows = [
                ["Dataset", "Identities", "Train / val / test", "Valid / selected", "Variation"],
                ["MOBIO", "150", "90 / 30 / 30", "1,799 / 1,800", "Sessions"],
                ["LFW", "125", "75 / 25 / 25", "1,500 / 1,500", "In-the-wild"],
                ["FEI", "200", "120 / 40 / 40", "2,378 / 2,400", "Pose / expression"],
                ["SCface", "130", "78 / 26 / 26", "2,851 / 2,860", "Camera / distance"],
            ]
            for row_index, row in enumerate(rows):
                for column, value in enumerate(row):
                    add_text(slide, fig, value, [0.6, 2.4, 4.15, 7.2, 9.9][column], 1.95 + row_index * 0.48,
                             [1.7, 1.7, 2.9, 2.6, 2.8][column], 0.45, size=14, bold=row_index == 0)
            add_text(slide, fig,
                     "SCface: mugshot-to-surveillance linkage under capture-domain shift.\n"
                     "Mugshot gallery / surveillance exposures; all 130 identities remain eligible after 9 detection failures.\n"
                     "Unprotected SCface top-1: 84.375% over 544 test probes; chance: 1/26 = 3.846%.\n"
                     f"Evidence: {len(table)} key-pool conditions / {table['source_file'].nunique()} studies; "
                     f"{len(pilots)} one-seed model endpoints / 24 scheme pilot cells.\n"
                     f"Earlier: {len(model_runs)} endpoints; 144 pool-study fits + 72 baseline evaluations; new: 672 MOBIO/SCface endpoints.",
                     0.6, 4.55, 12.1, 1.95, size=15)
        elif number == 2:
            add_text(slide, fig, "System overview: fixed encoder, hidden-key protection, same-identity exposure sets and held-out gallery linkage.",
                     0.6, 6.58, 12.1, 0.4, size=12)
        elif number == 3:
            add_text(slide, fig, "Mean pooling is the primary endpoint; DeepSets is secondary. Target is the normalized mean of exposed embeddings, not the gallery.",
                     0.6, 6.58, 12.1, 0.4, size=12)
        elif number == 4:
            add_text(slide, fig,
                     f"SCface: pool 3 rises from {100 * scface.loc['3', 'one_record_top1_mean']:.2f}% to "
                     f"{100 * scface.loc['3', 'top1_mean']:.2f}% top-1; fresh keys give {100 * scface.loc['fresh', 'top1_mean']:.2f}%.\n"
                     "SCface pools 1/2/3 pass the all-seed interval criterion; 4/5/7/10 fail. Failure is not equivalence.",
                     0.6, 1.05, 12.1, 0.65, size=14)
            add_text(slide, fig, f"{len(table)} conditions / {table['source_file'].nunique()} source-separated, three-seed studies. Grey: unavailable. *: interval failure. Pilots excluded.",
                     0.6, 6.58, 12.1, 0.4, size=12)
        elif number == 5:
            add_text(slide, fig,
                     "MOBIO: slot labels change DeepSets by -0.83 to +4.44 points; shuffled records collapse to 3.33% chance.\n"
                     "Same-image fresh-key sets also give 3.33%. Projection-sharing effects depend on construction and partition.",
                     0.6, 1.05, 12.1, 0.65, size=14)
            add_text(slide, fig, "Sources: experiments/mobio_mechanism_controls and mobio_correlation_controls. Three model seeds; chance 1/30.",
                     0.6, 6.58, 12.1, 0.4, size=12)
        elif number == 6:
            add_text(slide, fig, f"{len(model_runs)} trained endpoints; fresh, shared and pool-4 keys; 3 model seeds x 2 identity partitions.\n"
                     f"All {len(primary)} planned pool-4 amplification tests pass Holm correction (maximum adjusted p = {primary['holm_p'].max():.3f}).",
                     0.6, 1.05, 12.1, 0.65, size=14)
            add_text(slide, fig, "Matched 120-epoch caps. Crossed-bootstrap intervals include model-seed and identity resampling within each partition.",
                     0.6, 6.58, 12.1, 0.4, size=12)
        elif number == 7:
            add_text(slide, fig,
                     f"Fresh PolyProtect native identification exceeds the gallery-label null in {len(native_controls)} / {len(native_controls)} tests.\n"
                     f"Maximum Holm-adjusted p = {native_controls['holm_p'].max():.3f}; three fresh-key seeds per dataset and identity partition.",
                     0.6, 1.05, 12.1, 0.65, size=14)
            add_text(slide, fig, "Earlier synthetic stress: IoM-GRP unchanged, PolyProtect scale-sensitive. Genuine raw-input controls follow on page 12.",
                     0.6, 6.58, 12.1, 0.4, size=12)
        elif number == 8:
            values = " / ".join(f"{100 * fresh_native.loc[dataset, 'native_top1']:.2f}%" for dataset in ("MOBIO", "FEI", "SCface"))
            add_text(slide, fig, f"Earlier one-seed pilots: {len(pilots)} model endpoints on MOBIO / FEI / SCface, including pool 8.\n"
                     f"Separate native fresh-key PolyProtect diagnostic: {values} top-1, respectively.",
                     0.6, 1.05, 12.1, 0.65, size=14)
            add_text(slide, fig, "Pilot intervals condition on one seed and partition; not confirmation. New replicated MOBIO/FEI results are on page 6.",
                     0.6, 6.58, 12.1, 0.4, size=12)
        elif number == 9:
            add_text(slide, fig,
                     "WHAT PRIOR WORK ALREADY SHOWED\n"
                     "Record-multiplicity attacks: several templates plus known parameters can enable recovery.\n"
                     "PolyProtect (2022): 1-10-record inversion; residual linkage under naive random parameters;\n"
                     "stricter parameter selection improves unlinkability. Our native result is not a first discovery.\n"
                     "benchmark_cb: recognition, score-based unlinkability and information estimates across schemes.\n"
                     "FaceLinkGen v3: adaptive identity distillation from paired data with hidden per-query randomness.\n\n"
                     "Maximal-linkability work: joint similarity scores, multiple keys and composition limits (thesis 5.4).\n\n"
                     "WHAT OUR CONTROLLED EXPERIMENT ADDS\n"
                     "Different-image set aggregation with hidden fresh/shared/pool keys and disjoint training identities.\n"
                     "A measured one-to-ten benefit conditional on transform reuse, plus mechanism and failure controls.\n"
                     "Independent pool draws and prediction averaging expose pool sensitivity and a competitive baseline.\n\n"
                     "NOT CLAIMED\n"
                     "Multiplicity, identity distillation and IoM scale invariance are not new.\n"
                     "No head-to-head accuracy win, exhaustive priority claim or break of stricter PolyProtect policies.",
                     0.6, 1.15, 12.1, 5.45, size=13)
            add_text(slide, fig, "Full-text evidence: author thesis 5.4, linked to Access 2024; publisher PDF blocked. See docs/literature/closest_work_2026-09-18.md.",
                     0.6, 6.65, 12.1, 0.3, size=10)
        elif number == 10:
            add_text(slide, fig,
                     "SAME-PERSON RECORDS\n"
                     "An account or session pseudonym could group retained records without revealing real identity.\n"
                     "Stable grouping and retention are assumed; their prevalence in deployed products is not measured.\n\n"
                     "PAIRED TRAINING EXAMPLES\n"
                     "An authorized enrollment/query interface or a provider seeing inputs and outputs could supply pairs.\n"
                     "Crucially, training must use the SAME realized hidden pool as the target records.\n"
                     "A proxy with unrelated keys is insufficient. Identity splits are disjoint; recurring keys are not.\n\n"
                     "REFERENCE GALLERY\n"
                     "Lawful public or consented reference images; gallery images are excluded from exposure sets.\n"
                     "Only 25-40 candidate people, with guaranteed target membership: a small closed-set experiment.\n\n"
                     "BOUNDARIES\n"
                     "No open-set search, internet-scale distractors, unknown grouping or cross-provider transfer test.\n"
                     "Fresh independent transforms can also harm legitimate matching; they are not a tested defense.",
                     0.6, 1.15, 12.1, 5.45, size=15)
        elif number == 11:
            add_text(slide, fig,
                     "FORMULA AND PARAMETER CHECKS\n"
                     "PolyProtect: separate scalar windows, zero padding, coefficients and powers match local outputs.\n"
                     "IoM-GRP: explicit grouped dot-product/argmax reference agrees; paper dimensions are tested.\n"
                     "Original PolyProtect uses different encoders and offers a stricter parameter-selection policy.\n\n"
                     "NATIVE MATCHING AUDIT: 48 MATCHED CELLS\n"
                     "Scalar PolyProtect and production outputs: zero difference after float32 casting.\n"
                     "Matrix and SciPy cosine scores: maximum difference 1.33e-15; identical predictions.\n"
                     "Zero gallery/probe overlap, score ties, reversed-gallery changes or duplicate fresh keys.\n"
                     "The unit-input results reproduce the earlier three-key native averages.\n\n"
                     "RAW EXTRACTION CHECK\n"
                     "All 4,177 MOBIO/FEI records re-extracted; normalized discrepancy at most 2.98e-8.\n"
                     "Natural raw/unit scaling changes zero of 38,400 sampled IoM codes.\n\n"
                     "Separate computational checks, not independent human review or official-code certification.",
                     0.6, 1.15, 12.1, 5.45, size=15)
        elif number == 12:
            add_text(slide, fig, "Real raw embeddings, shuffled norms and a fixed training-median radius; same keys in every arm.\n"
                     "Raw matching rises to 15.58-16.66%, but shuffled norms perform similarly; identity-specific norm leakage is not established.",
                     0.6, 1.05, 12.1, 0.65, size=14)
            add_text(slide, fig, "Raw-unit gain survives only on FEI. This page is native matching; separate learned raw-input results are on page 18.",
                     0.6, 6.58, 12.1, 0.4, size=12)
        elif number == 13:
            add_text(slide, fig, "Post-hoc two-sided family of 48; original primary eight-test family remains unchanged.\n"
                     "Eight pool-4 mean gains survive; four shared-key PolyProtect DeepSets losses also survive (adjusted p = 0.0192).",
                     0.6, 1.05, 12.1, 0.65, size=14)
            add_text(slide, fig, "Seed ranges and leave-one-seed-out gains are published. Negative results constrain the claim: more records are not always better.",
                     0.6, 6.58, 12.1, 0.4, size=12)
        elif number == 14:
            add_text(slide, fig,
                     "TRANSFORM REUSE AND RECORD MULTIPLICITY\n"
                     "Recurring hidden transforms support substantial identity linkage from multiple protected records.\n"
                     "The transition depends on dataset, partition and realized pool, not a universal pool-size threshold.\n"
                     "SCface pool 3: 5.29% single-record versus 33.17% ten-record top-1; chance is 3.846%.\n\n"
                     "BOUNDARIES AND FAILURES\n"
                     "Shuffled-identity controls constrain the mechanism; same-pool training access remains essential.\n"
                     "Shared-key PolyProtect DeepSets loses 15.42-24.86 points: aggregation is not universally beneficial.\n"
                     "Native matching survives separate checks, but natural identity-specific norm leakage is not established.\n\n"
                     "STATISTICAL AND THREAT-MODEL SCOPE\n"
                     "BioHash spans four datasets; scheme follow-ups cover MOBIO/FEI and now MOBIO/SCface.\n"
                     f"New scheme follow-up: {len(model_runs)} endpoints; all {len(primary)} primary contrasts pass Holm correction.\n"
                     "New three-pool study: IoM gains persist; PolyProtect is sensitive to the pool draw.\n"
                     "All direct prediction-mean baseline intervals include zero: no input-pooling superiority claim.\n"
                     "Fresh-key learned intervals include chance; they do not establish equivalence or general unlinkability.\n\n"
                     "Scope: only three new pools, fixed set seed, overlapping assignments, one encoder, small galleries, paired access.\n"
                     "Three extended-study source snapshots remain unresolved locally; independent human review remains open.",
                     0.6, 1.15, 12.1, 5.45, size=13)
        elif number == 15:
            add_text(slide, fig, "All four datasets are shown, without mixing historical studies, one-seed pilots and matched follow-ups.\n"
                     "SCface now has extended training and controls; LFW and FEI do not have the same complete new matrix.",
                     0.6, 1.05, 12.1, 0.65, size=14)
            add_text(slide, fig, "Evidence coverage, not a performance ranking. LFW remains historical BioHash evidence; missing studies require experiments.",
                     0.6, 6.58, 12.1, 0.4, size=12)
        elif number == 16:
            add_text(slide, fig, "24 cells; 3 new pool draws x 3 model seeds x 2 partitions; 144 fits plus 72 prediction-mean endpoints.\n"
                     "Matched gallery and records; prediction mean reuses single-record training, not the set-training objective.\n"
                     "Pool seeds vary transforms and slot assignments jointly. Intervals are pointwise; only three pool draws.",
                     0.6, 1.05, 12.1, 0.8, size=12)
            add_text(slide, fig, "IoM per-pool gains: +21.56 to +42.36 pp. PolyProtect: -3.23 to +72.40 pp. Three pools do not establish universal robustness.",
                     0.6, 6.58, 12.1, 0.4, size=12)
        elif number == 17:
            add_text(slide, fig, "32 cells / 672 endpoints: 3 seeds x 2 splits; exposures 1/2/5/10; fresh and pools 1/4/8.\n"
                     "Eight pool-4 primary gains pass Holm correction (p = 0.004-0.034); all 56 fresh endpoint intervals include chance.",
                     0.6, 1.05, 12.1, 0.65, size=13)
            add_text(slide, fig, "SCface pool-4 gains: IoM +16.19/+34.29 pp; PolyProtect +5.93/+5.77 pp. One pool-8 draw cannot isolate a pool-size effect.",
                     0.6, 6.58, 12.1, 0.4, size=12)
        elif number == 18:
            add_text(slide, fig, "24 summaries / 72 fits; one saved partition; three model seeds; raw embeddings before protection.\n"
                     "Pool-4 mean top-1: PolyProtect 3.61% / 3.85%; IoM 91.11% / 67.63% on MOBIO / SCface.",
                     0.6, 1.05, 12.1, 0.65, size=14)
            add_text(slide, fig, "Descriptive, not a corrected null test. Raw inputs also change set targets; normalization alone is not isolated. No privacy claim.",
                     0.6, 6.58, 12.1, 0.4, size=12)
        elif number == 19:
            add_text(slide, fig, "Native protected-gallery matching, not learned reconstruction: three fixed key seeds per dataset.\n"
                     "MOBIO: 13.13% -> 16.26%; SCface: 10.44% -> 10.12%. Neither paired change survives correction.",
                     0.6, 1.05, 12.1, 0.65, size=14)
            add_text(slide, fig, "All four arms exceed the gallery-label null (Holm p = 0.0008). This does not certify or invalidate every stricter selection policy.",
                     0.6, 6.58, 12.1, 0.4, size=12)
        buffer = BytesIO()
        fig.savefig(buffer, format="pdf")
        plt.close(fig)
        buffer.seek(0)
        page = writer.add_page(PdfReader(buffer).pages[0])
        if number in figure_names:
            source = PdfReader(FIGURES / f"{figure_names[number]}.pdf").pages[0]
            top = 1.05 if number in (2, 3) else 1.9
            available_height = 5.3 if number in (2, 3) else 4.5
            scale = min(12.1 * 72 / float(source.mediabox.width), available_height * 72 / float(source.mediabox.height))
            left = 0.6 * 72 + (12.1 * 72 - float(source.mediabox.width) * scale) / 2
            bottom = (HEIGHT - top - available_height) * 72 + (available_height * 72 - float(source.mediabox.height) * scale) / 2
            page.merge_transformed_page(source, Transformation().scale(scale).translate(left, bottom))
    writer.add_metadata({"/Title": "September Dataset Update - 19 September 2026",
                         "/Subject": "Identity-disjoint multi-record linkage: architecture, results and study scope"})
    destination.parent.mkdir(parents=True, exist_ok=True)
    temporary = destination.with_suffix(".tmp.pdf")
    writer.write(temporary)
    temporary.replace(destination)
    print(f"{len(titles)}-page vector dataset update -> {destination}")


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--out", type=Path, default=ROOT / "reports/slides")
    parser.add_argument("--dataset-update", type=Path, default=ROOT / "reports/Sept_Dataset_Update.pdf")
    args = parser.parse_args()
    build_presentation(args.out)
    build_dataset_update(args.dataset_update)


if __name__ == "__main__":
    main()
